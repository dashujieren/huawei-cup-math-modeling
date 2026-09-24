"""第四问基线的原子块、贪心、资源峰值与文件输出验收。"""

import csv
import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from q4_0_baseline import (
    CODE_DIR,
    Occupancy,
    Q4_COLUMNS,
    Relay,
    Transport,
    calculate,
    charge_time_s,
    latest_data_run,
    load_clean_data,
    load_q3,
    peak_occupancy,
    save_results,
    greedy_partition,
)


def _write(path: Path, columns: tuple[str, ...], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def test_charge_time_matches_two_stage_rule() -> None:
    assert charge_time_s(0.0, 1800) == pytest.approx(1800)
    assert charge_time_s(0.9, 1800) == pytest.approx(630)
    assert charge_time_s(1.0, 1800) == pytest.approx(0)
    assert charge_time_s(0.8, 1800) == pytest.approx(
        1800 * (0.65 * 0.1 / 0.9 + 0.35)
    )
    with pytest.raises(ValueError, match="SOC"):
        charge_time_s(-0.1, 1800)


def test_q4_columns_match_official_result_template() -> None:
    template = CODE_DIR.parent / "0_3_数据" / "结果提交模板.xlsx"
    workbook = load_workbook(template, read_only=True, data_only=True)
    try:
        actual = tuple(
            cell.value for cell in next(workbook["Q4_分区配置"].iter_rows(max_row=1))
        )[:len(Q4_COLUMNS)]
    finally:
        workbook.close()
    assert actual == Q4_COLUMNS


def test_half_open_resource_peak() -> None:
    intervals = [
        Occupancy("transport", "T1", "UAV_A", 0, 10),
        Occupancy("transport", "T2", "UAV_A", 10, 20),
    ]
    assert peak_occupancy(intervals) == 1
    assert peak_occupancy(intervals + [
        Occupancy("transport", "T3", "UAV_A", 9, 12)
    ]) == 2


def test_three_groups_require_three_atomic_blocks() -> None:
    blocks = [
        {"block_id": "B01", "zones": ["S001"], "box_count": 2},
        {"block_id": "B02", "zones": ["S002"], "box_count": 1},
    ]
    with pytest.raises(ValueError, match="返回问题三"):
        greedy_partition(blocks, 3)


def test_independent_groups_require_more_resources_than_global_peak() -> None:
    zones = {"S001", "S002", "S003"}
    clean = {
        "zones": zones,
        "boxes": {
            f"B{i}": {
                "zone_id": zone, "mass_kg": "1", "priority": "1",
                "medical_bool": "False", "is_first_batch": "False",
            }
            for i, zone in enumerate(sorted(zones), 1)
        },
        "types": {
            "A": {"energy_use_kwh": "4", "return_soc_min": "0.2"},
        },
        "battery_by_type": {"A": {"full_charge_s": "100"}},
        "relay_type": {
            "energy_use_kwh": "3", "return_soc_min": "0.2",
            "turnaround_s": "300",
        },
        "relay_pool": {"full_charge_s": "1800"},
        "stock": {
            "UAV_A": 1, "UAV_B": 0, "UAV_C": 0,
            "BAT_A": 1, "BAT_B": 0, "BAT_C": 0,
            "RELAY_UAV": 0, "RELAY_ENERGY": 0,
        },
    }
    tasks = {
        f"T{i}": Transport(f"T{i}", "A", (i - 1) * 20, (i - 1) * 20 + 10,
                           0, (zone,))
        for i, zone in enumerate(sorted(zones), 1)
    }
    result = calculate(clean, tasks, {}, {})
    assert result["global_peak"]["UAV_A"] == 1
    assert result["solutions"][2]["totals"]["UAV_A"] == 2
    assert result["solutions"][2]["deficit"]["UAV_A"] == 1
    assert result["solutions"][3]["totals"]["UAV_A"] == 3
    assert all(result["solutions"][3]["group_zones"].values())
    reused = {
        task_id: replace(task, uav_id="U01", battery_id="A-B01")
        for task_id, task in tasks.items()
    }
    reused["T2"] = replace(reused["T2"], start_s=5, return_s=15)
    with pytest.raises(ValueError, match="原编号资源占用冲突"):
        calculate(clean, reused, {}, {})


def test_q3_inputs_merge_relay_zones_and_export_template_rows(tmp_path: Path) -> None:
    try:
        data_run = latest_data_run()
    except FileNotFoundError:
        pytest.skip("需要先运行 0_0_run_all.py")
    clean = load_clean_data(data_run)
    q3_dir = tmp_path / "q3"
    q3_dir.mkdir()
    zones = sorted(clean["zones"])
    transport_rows = []
    delivery_rows = []
    comm_rows = []
    for index, zone in enumerate(zones, 1):
        task_id = f"T{index:03d}"
        start = (index - 1) * 800
        transport_rows.append({
            "架次编号": task_id, "机型编号": "A",
            "无人机编号": "U01", "电池编号": "A-B01",
            "开始时刻（s）": start, "访问服务区顺序": zone,
            "返回O01时刻（s）": start + 500, "架次能耗（kWh）": 0.1,
        })
        for box in clean["boxes"].values():
            if box["zone_id"] == zone:
                delivery_rows.append({
                    "货箱编号": box["box_id"], "架次编号": task_id,
                    "服务区编号": zone, "交付完成时刻（s）": start + 300,
                })
        comm_rows.append({
            "运输架次编号": task_id, "开始时刻（s）": start + 100,
            "结束时刻（s）": start + 200,
            "中继架次编号": "R001" if index in (1, 2) else "",
        })
    relay_rows = [
        {
            "中继架次编号": "R001", "开始时刻（s）": 0,
            "中继无人机编号": "R01", "能源组件编号": "R-E01",
            "建链完成时刻（s）": 50, "服务结束时刻（s）": 1100,
            "返回O01时刻（s）": 1200, "架次能耗（kWh）": 0.5,
        },
    ]
    _write(q3_dir / "Q3_运输架次.csv",
           ("架次编号", "无人机编号", "机型编号", "电池编号",
            "开始时刻（s）",
            "访问服务区顺序", "返回O01时刻（s）", "架次能耗（kWh）"),
           transport_rows)
    _write(q3_dir / "Q3_逐箱交付.csv",
           ("货箱编号", "架次编号", "服务区编号", "交付完成时刻（s）"),
           delivery_rows)
    _write(q3_dir / "Q3_中继架次.csv",
           ("中继架次编号", "中继无人机编号", "能源组件编号",
            "开始时刻（s）", "建链完成时刻（s）",
            "服务结束时刻（s）", "返回O01时刻（s）", "架次能耗（kWh）"),
           relay_rows)
    _write(q3_dir / "Q3_通信保障.csv",
           ("运输架次编号", "开始时刻（s）", "结束时刻（s）",
            "中继架次编号"), comm_rows)
    (q3_dir / "Q3_READY.txt").write_text("synthetic fixture only", encoding="utf-8")
    transports, relays, relation = load_q3(q3_dir, clean)
    result = calculate(clean, transports, relays, relation)
    assert len(result["blocks"]) == 14
    assert result["solutions"][2]["relay_group"]["R001"] == (
        result["solutions"][2]["transport_group"]["T001"]
    )
    assert result["solutions"][3]["relay_group"]["R001"] == (
        result["solutions"][3]["transport_group"]["T002"]
    )
    output_dir = tmp_path / "q4"
    output_dir.mkdir()
    save_results(output_dir, q3_dir, data_run, clean, transports, relays, relation,
                 result)
    with (output_dir / "Q4_分区配置_基线.csv").open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 5
    assert {row["K（2或3）"] for row in rows} == {"2", "3"}
    assert (output_dir / "Q4_BASE_READY.txt").is_file()
    assert (output_dir / "Q4_原编号保守配置_基线.csv").is_file()
    summary = json.loads((output_dir / "Q4_运行摘要.json").read_text(encoding="utf-8"))
    assert summary["transport_sorties"] == 15
    assert summary["atomic_block_count"] == 14
    assert summary["original_ids_complete"]
    assert summary["by_k"]["3"]["original_id_copy_demand"]["UAV_A"] == 3
