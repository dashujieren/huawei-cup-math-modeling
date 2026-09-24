"""D 题问题二基线：同区组批，A/B/C 三种机型参与简单列表排程。

单次运行先写 Q2 官方格式表和验收表，再画实体机/电池及交付时限图。
这是截止优先的事件排程基线，不是跨区联合优化；超时会明确标记为失败。
"""

from __future__ import annotations

import argparse
import json
import math
import os
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from _0_pipeline import sha256, verify_ready
from q1_0_baseline import CODE_DIR, PROJECT, _read_inputs, find_latest_validated_run, sortie_metrics


TOL = 1e-7
SORTIE_COLUMNS = [
    "架次编号", "无人机编号", "机型编号", "电池编号", "开始时刻（s）",
    "访问服务区顺序", "返回O01时刻（s）", "架次能耗（kWh）",
]
DELIVERY_COLUMNS = ["货箱编号", "架次编号", "服务区编号", "交付完成时刻（s）"]


@dataclass
class UAV:
    uav_id: str
    type_id: str
    available_s: float = 0.0


@dataclass
class Battery:
    battery_id: str
    type_id: str
    full_charge_s: float
    available_s: float = 0.0


def optional_seconds(value: Any) -> float | None:
    if value is None or pd.isna(value) or str(value).strip() == "":
        return None
    seconds = float(value)
    if not math.isfinite(seconds) or seconds < 0:
        raise ValueError(f"非法时刻：{value}")
    return seconds


def required_seconds(value: Any, label: str) -> float:
    seconds = optional_seconds(value)
    if seconds is None:
        raise ValueError(f"缺少{label}")
    return seconds


def source_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if str(value).strip().lower() in {"true", "1", "是"}:
        return True
    if str(value).strip().lower() in {"false", "0", "否"}:
        return False
    raise ValueError(f"非法布尔值：{value}")


def charge_to_full_s(return_soc: float, full_charge_s: float) -> float:
    """官方两阶段模型：0–90% 用 65% 总时长，90–100% 用 35%。"""
    if not 0 <= return_soc <= 1 + TOL or full_charge_s <= 0:
        raise ValueError("电池 SOC 或等效充满时间超范围")
    soc = min(return_soc, 1.0)
    if soc < 0.9:
        return full_charge_s * (0.65 * (0.9 - soc) / 0.9 + 0.35)
    return full_charge_s * 0.35 * (1.0 - soc) / 0.1


def read_resources(data_run: Path) -> tuple[list[UAV], list[Battery]]:
    clean = data_run / "clean"
    uav_rows = pd.read_csv(clean / "transport_uavs.csv", dtype=str).to_dict("records")
    battery_rows = pd.read_csv(clean / "battery_units.csv", dtype=str).to_dict("records")
    pool_rows = pd.read_csv(clean / "battery_pool.csv", dtype=str).to_dict("records")
    if len(uav_rows) != 8 or len(battery_rows) != 14:
        raise ValueError("问题二要求 8 架实体运输机和总计 14 组电池")
    uavs = [UAV(str(row["uav_id"]), str(row["type_id"])) for row in uav_rows]
    if len({u.uav_id for u in uavs}) != 8 or any(row["initial_site"] != "O01" for row in uav_rows):
        raise ValueError("运输机编号不唯一或初始位置不是 O01")
    pool: dict[str, tuple[int, float]] = {}
    for row in pool_rows:
        type_id = str(row["type_id"])
        if type_id in pool:
            raise ValueError(f"重复的电池机型：{type_id}")
        pool[type_id] = (int(row["stock_qty"]), float(row["full_charge_s"]))
    if {kind: qty for kind, (qty, _) in pool.items()} != {"A": 6, "B": 4, "C": 4}:
        raise ValueError("A/B/C 电池库存与附件 6/4/4 不一致")
    batteries = []
    for row in battery_rows:
        type_id = str(row["type_id"])
        if type_id not in pool or not math.isclose(float(row["initial_soc"]), 1.0, abs_tol=TOL):
            raise ValueError("电池机型或初始 100% SOC 不符合附件")
        batteries.append(Battery(str(row["battery_id"]), type_id, pool[type_id][1]))
    if len({b.battery_id for b in batteries}) != 14:
        raise ValueError("电池编号不唯一")
    if Counter(u.type_id for u in uavs) != Counter({"A": 4, "B": 2, "C": 2}):
        raise ValueError("实体机 A/B/C 数量与附件 4/2/2 不一致")
    if Counter(b.type_id for b in batteries) != Counter({"A": 6, "B": 4, "C": 4}):
        raise ValueError("电池明细与库存不一致，不能另加初装电池")
    return sorted(uavs, key=lambda u: u.uav_id), sorted(batteries, key=lambda b: b.battery_id)


def ordered_boxes(batch: dict[str, Any], boxes_by_id: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    boxes = [boxes_by_id[box_id] for box_id in batch["box_ids"]]
    # 同一站点先交付硬时限紧的箱，接着按期望时刻与优先系数排序。
    return sorted(boxes, key=lambda b: (
        optional_seconds(b["hard_deadline_s"]) if optional_seconds(b["hard_deadline_s"]) is not None else math.inf,
        required_seconds(b["expected_s"], f"{b['box_id']} 的期望时刻"),
        -float(b["priority"]), str(b["box_id"]),
    ))


def make_tasks(data_run: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]], dict[str, Any]]:
    boxes, models, arcs = _read_inputs(data_run)
    boxes_by_id = {str(box["box_id"]): box for box in boxes}
    if len(boxes_by_id) != 80:
        raise ValueError("货箱编号不唯一")
    tasks: list[dict[str, Any]] = []
    seen: Counter[str] = Counter()
    by_zone: dict[str, list[dict[str, Any]]] = {}
    for box in boxes:
        by_zone.setdefault(str(box["zone_id"]), []).append(box)

    def physical_for(kind: str, zone: str, group: list[dict[str, Any]]) -> dict[str, Any] | None:
        model = models[kind]
        mass = sum(float(box["mass_kg"]) for box in group)
        volume = sum(float(box["volume_m3"]) for box in group)
        # sortie_metrics 在返回 feasible=False 前会先计算航段能耗；超载必须在此拦住。
        if (mass > float(model["max_payload_kg"]) + TOL or
                volume > float(model["capacity_m3"]) + TOL):
            return None
        ordered = ordered_boxes({"box_ids": [str(b["box_id"]) for b in group]}, boxes_by_id)
        physical = sortie_metrics(model, arcs[("O01", zone)], arcs[(zone, "O01")], ordered)
        if not physical["feasible"]:
            return None
        return {**physical,
                "handoff_base_s": float(model["handoff_base_s"]),
                "handoff_per_box_s": float(model["handoff_per_box_s"]),
                "return_soc_min": float(model["return_soc_min"])}

    def append_task(zone: str, group: list[dict[str, Any]], tier: str) -> None:
        batch_id = f"Q{len(tasks) + 1:03d}"
        delivery_order = ordered_boxes({"box_ids": [str(b["box_id"]) for b in group]}, boxes_by_id)
        options = {kind: physical for kind in sorted(models)
                   if (physical := physical_for(kind, zone, delivery_order)) is not None}
        if not options:
            raise ValueError(f"{batch_id} 对 A/B/C 均不满足载重、体积或返航余量")
        seen.update(str(box["box_id"]) for box in delivery_order)
        for box in delivery_order:
            if source_bool(box["medical_bool"]) or source_bool(box["is_first_batch"]):
                if optional_seconds(box["hard_deadline_s"]) is None:
                    raise ValueError(f"{box['box_id']} 缺少医疗或首批硬时限")
        deadlines = [optional_seconds(box["hard_deadline_s"]) for box in delivery_order]
        tasks.append({
            "batch_id": batch_id, "zone_id": zone, "tier": tier,
            "boxes": delivery_order, "options": options,
            "nearest_hard_s": min((d for d in deadlines if d is not None), default=math.inf),
            "nearest_expected_s": min(required_seconds(b["expected_s"], "期望时刻") for b in delivery_order),
        })

    for zone in sorted(by_zone):
        hard = [b for b in by_zone[zone] if optional_seconds(b["hard_deadline_s"]) is not None]
        routine = sorted((b for b in by_zone[zone] if optional_seconds(b["hard_deadline_s"]) is None),
                         key=lambda b: (-float(b["mass_kg"]), -float(b["volume_m3"]), str(b["box_id"])))
        if hard:
            append_task(zone, hard, "hard_same_zone")
        bins: list[list[dict[str, Any]]] = []
        for box in routine:
            for group in bins:
                if physical_for("C", zone, [*group, box]) is not None:
                    group.append(box)
                    break
            else:
                if physical_for("C", zone, [box]) is None:
                    raise ValueError(f"{box['box_id']} 无法由 C 型机单独运送")
                bins.append([box])
        for group in bins:
            append_task(zone, group, "routine_same_zone")
    if seen != Counter(boxes_by_id.keys()):
        raise ValueError("基线未将 80 个货箱恰好分配一次")
    return sorted(tasks, key=lambda t: (t["nearest_hard_s"], t["nearest_expected_s"], t["batch_id"])), boxes_by_id, {
        "hard_group_tasks": sum(t["tier"] == "hard_same_zone" for t in tasks),
        "routine_tasks": sum(t["tier"] == "routine_same_zone" for t in tasks),
    }


def schedule_baseline(tasks: list[dict[str, Any]], uavs: list[UAV],
                      batteries: list[Battery]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    sorties: list[dict[str, Any]] = []
    deliveries: list[dict[str, Any]] = []
    used_types: set[str] = set()
    for task in tasks:
        choices = []
        for kind, physical in task["options"].items():
            for uav in uavs:
                if uav.type_id != kind:
                    continue
                for battery in batteries:
                    if battery.type_id != kind:
                        continue
                    start = max(uav.available_s, battery.available_s)
                    arrival = (start + physical["prep_s"] + physical["loading_s"] +
                               physical["flight_out_s"])
                    hard_late = []
                    soft_late = 0.0
                    for index, box in enumerate(task["boxes"], 1):
                        complete = (arrival + physical["handoff_base_s"] +
                                    index * physical["handoff_per_box_s"])
                        hard = optional_seconds(box["hard_deadline_s"])
                        if hard is not None and complete > hard + TOL:
                            hard_late.append(complete - hard)
                        if hard is None:
                            soft_late += float(box["priority"]) * max(
                                0.0, complete - required_seconds(box["expected_s"], str(box["box_id"])))
                    # 三种机型都作为候选；同样满足硬时限时优先启用尚未使用的机型。
                    rank = (len(hard_late), sum(hard_late), kind in used_types,
                            soft_late, start + physical["operation_s"],
                            physical["energy_kwh"], kind, uav.uav_id, battery.battery_id)
                    choices.append((rank, kind, physical, start, uav, battery))
        if not choices:
            raise ValueError(f"{task['batch_id']} 无可匹配机型、实体机或电池")
        _, kind, physical, start_s, uav, battery = min(choices, key=lambda x: x[0])
        prep_end_s = start_s + physical["prep_s"]
        launch_s = prep_end_s + physical["loading_s"]
        arrival_s = launch_s + physical["flight_out_s"]
        handoff_end_s = arrival_s + physical["handoff_s"]
        return_s = handoff_end_s + physical["flight_back_s"]
        soc = float(physical["return_soc"])
        charge_s = charge_to_full_s(soc, battery.full_charge_s)
        charge_end_s = return_s + charge_s
        if not math.isclose(return_s, start_s + physical["operation_s"], rel_tol=0, abs_tol=TOL):
            raise AssertionError("排程的返回时刻与问题一作业时间不一致")
        sorties.append({
            "batch_id": task["batch_id"], "zone_id": task["zone_id"], "type_id": kind,
            "uav_id": uav.uav_id, "battery_id": battery.battery_id,
            "box_ids": [str(box["box_id"]) for box in task["boxes"]],
            "start_s": start_s, "prep_end_s": prep_end_s, "launch_s": launch_s,
            "arrival_s": arrival_s, "handoff_end_s": handoff_end_s, "return_s": return_s,
            "charge_start_s": return_s, "charge_end_s": charge_end_s, "charge_s": charge_s,
            "mass_kg": physical["mass_kg"], "volume_m3": physical["volume_m3"],
            "energy_kwh": physical["energy_kwh"], "return_soc": soc,
        })
        for index, box in enumerate(task["boxes"], 1):
            # 基础交接只收取一次；逐箱交接时间由对应机型参数给定。
            complete_s = (arrival_s + physical["handoff_base_s"] +
                          index * physical["handoff_per_box_s"])
            hard_s = optional_seconds(box["hard_deadline_s"])
            expected_s = required_seconds(box["expected_s"], f"{box['box_id']} 的期望时刻")
            deliveries.append({
                "box_id": str(box["box_id"]), "batch_id": task["batch_id"],
                "zone_id": task["zone_id"], "sequence": index, "complete_s": complete_s,
                "hard_deadline_s": hard_s, "expected_s": expected_s,
                "priority": float(box["priority"]),
                "hard_slack_s": None if hard_s is None else hard_s - complete_s,
                "soft_lateness_s": max(0.0, complete_s - expected_s),
            })
        uav.available_s = return_s
        battery.available_s = charge_end_s
        used_types.add(kind)
    return sorted(sorties, key=lambda s: (s["start_s"], s["batch_id"])), sorted(deliveries, key=lambda d: d["box_id"])


def check_plan(sorties: list[dict[str, Any]], deliveries: list[dict[str, Any]],
               tasks: list[dict[str, Any]], uavs: list[UAV], batteries: list[Battery],
               boxes_by_id: dict[str, dict[str, Any]]) -> list[dict[str, str]]:
    checks: list[dict[str, str]] = []

    def add(check_id: str, good: bool, detail: str = "") -> None:
        checks.append({"check_id": check_id, "status": "PASS" if good else "FAIL", "detail": detail})

    add("all_80_boxes_once", Counter(d["box_id"] for d in deliveries) == Counter(boxes_by_id.keys()),
        f"deliveries={len(deliveries)}, unique={len({d['box_id'] for d in deliveries})}")
    add("all_baseline_batches_once", Counter(s["batch_id"] for s in sorties) ==
        Counter(t["batch_id"] for t in tasks), f"sorties={len(sorties)}")
    add("A_B_C_all_used", {s["type_id"] for s in sorties} == {"A", "B", "C"},
        f"type_sorties={dict(Counter(s['type_id'] for s in sorties))}")
    uav_types = {u.uav_id: u.type_id for u in uavs}
    battery_types = {b.battery_id: b.type_id for b in batteries}
    battery_full_s = {b.battery_id: b.full_charge_s for b in batteries}
    tasks_by_id = {t["batch_id"]: t for t in tasks}
    sorties_by_id = {s["batch_id"]: s for s in sorties}
    for sortie in sorties:
        batch_id = sortie["batch_id"]
        add(f"{batch_id}:same_type", uav_types[sortie["uav_id"]] ==
            battery_types[sortie["battery_id"]] == sortie["type_id"])
        expected = tasks_by_id[batch_id]["options"].get(sortie["type_id"])
        add(f"{batch_id}:assigned_type_feasible", expected is not None)
        if expected is None:
            continue
        add(f"{batch_id}:return_soc",
            sortie["return_soc"] >= expected["return_soc_min"] - TOL,
            f"SOC={sortie['return_soc']:.9f}, minimum={expected['return_soc_min']:.9f}")
        add(f"{batch_id}:time_order", 0 <= sortie["start_s"] <= sortie["launch_s"] <=
            sortie["arrival_s"] <= sortie["handoff_end_s"] <= sortie["return_s"] <=
            sortie["charge_end_s"])
        add(f"{batch_id}:one_zone", all(boxes_by_id[box_id]["zone_id"] == sortie["zone_id"]
                                       for box_id in sortie["box_ids"]))
        add(f"{batch_id}:physical_replay", math.isclose(sortie["energy_kwh"],
            expected["energy_kwh"], rel_tol=0, abs_tol=TOL) and
            math.isclose(sortie["return_s"] - sortie["start_s"],
                         expected["operation_s"], rel_tol=0, abs_tol=TOL))
        expected_charge_s = charge_to_full_s(
            sortie["return_soc"], battery_full_s[sortie["battery_id"]])
        add(f"{batch_id}:charging_replay", math.isclose(
            sortie["charge_end_s"] - sortie["return_s"], expected_charge_s,
            rel_tol=0, abs_tol=TOL))
    for delivery in deliveries:
        task = tasks_by_id[delivery["batch_id"]]
        sortie = sorties_by_id[delivery["batch_id"]]
        physical = task["options"][sortie["type_id"]]
        expected_delivery_s = (sortie["arrival_s"] + physical["handoff_base_s"] +
                               delivery["sequence"] * physical["handoff_per_box_s"])
        add(f"{delivery['box_id']}:delivery_replay", math.isclose(
            delivery["complete_s"], expected_delivery_s, rel_tol=0, abs_tol=TOL) and
            delivery["complete_s"] <= sortie["handoff_end_s"] + TOL)
        hard = delivery["hard_deadline_s"]
        if hard is not None:
            add(f"{delivery['box_id']}:hard_deadline", delivery["complete_s"] <= hard + TOL,
                f"delivered={delivery['complete_s']:.3f}s, deadline={hard:.3f}s, "
                f"slack={delivery['hard_slack_s']:.3f}s")
    for uav in uavs:
        uses = sorted((s for s in sorties if s["uav_id"] == uav.uav_id), key=lambda s: s["start_s"])
        add(f"{uav.uav_id}:no_overlap", all(a["return_s"] <= b["start_s"] + TOL
                                            for a, b in zip(uses, uses[1:])))
    for battery in batteries:
        uses = sorted((s for s in sorties if s["battery_id"] == battery.battery_id),
                      key=lambda s: s["start_s"])
        add(f"{battery.battery_id}:charged_before_reuse",
            all(a["charge_end_s"] <= b["start_s"] + TOL for a, b in zip(uses, uses[1:])))
    return checks


def validate_template() -> None:
    template = load_workbook(PROJECT / "0_3_数据" / "结果提交模板.xlsx", read_only=True, data_only=True)
    try:
        for sheet_name, expected in (("Q2_运输架次", SORTIE_COLUMNS),
                                     ("Q2_逐箱交付", DELIVERY_COLUMNS)):
            actual = [template[sheet_name].cell(1, col).value for col in range(1, len(expected) + 1)]
            if actual != expected:
                raise ValueError(f"官方模板 {sheet_name} 列名变化，停止导出：{actual}")
    finally:
        template.close()


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    pd.DataFrame(rows, columns=columns).to_csv(path, index=False, encoding="utf-8-sig", float_format="%.12g")


def save_delivery_figures(figure_dir: Path, deliveries: list[dict[str, Any]], plt: Any) -> None:
    """用逐箱交接完成时刻展示进度和硬时限，不使用到站时刻。"""
    ordered = sorted(deliveries, key=lambda row: (row["complete_s"], row["box_id"]))
    hard = [row for row in ordered if row["hard_deadline_s"] is not None]
    if not ordered or not hard:
        raise ValueError("缺少逐箱交付记录或硬时限记录，无法绘制问题二进度图")

    fig, ax = plt.subplots(figsize=(9, 5))
    hours = [0.0, *(row["complete_s"] / 3600 for row in ordered)]
    ax.step(hours, range(len(hours)), where="post", color="#2879B8", lw=2,
            label="全部货箱")
    hard_hours = [0.0, *(row["complete_s"] / 3600 for row in hard)]
    ax.step(hard_hours, range(len(hard_hours)), where="post", color="#D98537",
            lw=1.8, label="硬时限货箱")
    for index, deadline in enumerate(sorted({row["hard_deadline_s"] for row in hard})):
        ax.axvline(deadline / 3600, ls="--", lw=0.9, color="#777777",
                   alpha=0.75, label="硬截止" if index == 0 else None)
    ax.set(xlabel="从开始至交接完成（小时）", ylabel="累计完成交接箱数",
           ylim=(0, len(ordered) + 3))
    ax.grid(axis="y", alpha=0.22)
    ax.legend(frameon=False, ncol=3)
    fig.tight_layout()
    for extension in ("png", "pdf"):
        fig.savefig(figure_dir / f"Q2_逐箱交付进度.{extension}", dpi=220)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.5, 6))
    on_time = [row for row in hard if row["hard_slack_s"] >= -TOL]
    late = [row for row in hard if row["hard_slack_s"] < -TOL]
    for subset, color, label in ((on_time, "#2879B8", "按时交接"),
                                 (late, "#C64F4F", "超时交接")):
        if subset:
            ax.scatter([row["hard_deadline_s"] / 3600 for row in subset],
                       [row["complete_s"] / 3600 for row in subset],
                       s=38, color=color, alpha=0.78, label=f"{label}（{len(subset)}箱）")
    limit = max(max(row["hard_deadline_s"], row["complete_s"]) for row in hard) / 3600
    ax.plot([0, limit * 1.04], [0, limit * 1.04], ls="--", lw=1,
            color="#4D4D4D", label="交接完成 = 硬截止")
    ax.set(xlabel="硬截止时刻（小时）", ylabel="交接完成时刻（小时）",
           xlim=(0, limit * 1.04), ylim=(0, limit * 1.04))
    ax.set_aspect("equal", adjustable="box")
    ax.grid(alpha=0.2)
    ax.legend(frameon=False, loc="upper left")
    fig.tight_layout()
    for extension in ("png", "pdf"):
        fig.savefig(figure_dir / f"Q2_交付硬时限对照.{extension}", dpi=220)
    plt.close(fig)


def save_route_figure(figure_dir: Path, data_run: Path,
                      sorties: list[dict[str, Any]], plt: Any) -> None:
    """在同一投影坐标上画各机型实际飞行顺序，重复路线合并显示。"""
    from matplotlib.patches import Patch

    frame = pd.read_csv(data_run / "clean" / "nodes.csv", dtype={"node_id": str})
    nodes = {str(row["node_id"]): (float(row["x_m"]) / 1000,
                                   float(row["y_m"]) / 1000)
             for row in frame.to_dict("records")}
    colors = {"A": "#3A80C1", "B": "#27A275", "C": "#DB873F"}
    routes: Counter[tuple[str, tuple[str, ...]]] = Counter()
    for sortie in sorties:
        if "route" in sortie:
            zones = tuple(sortie["route"].zones)
        elif "visit_order" in sortie:
            zones = tuple(str(sortie["visit_order"]).split(">"))
        else:
            zones = (str(sortie["zone_id"]),)
        routes[(sortie["type_id"], zones)] += 1

    fig, ax = plt.subplots(figsize=(9, 7))
    for (kind, zones), count in sorted(routes.items()):
        sequence = ("O01", *zones, "O01")
        for index, (start, end) in enumerate(zip(sequence, sequence[1:]), 1):
            x0, y0 = nodes[start]
            x1, y1 = nodes[end]
            dx, dy = x1 - x0, y1 - y0
            length = math.hypot(dx, dy)
            if length <= TOL:
                continue
            # 相反方向的箭头落在航线两侧，往返重合时仍能辨认。
            shift_x, shift_y = -dy / length * 0.08, dx / length * 0.08
            x0, y0, x1, y1 = x0 + shift_x, y0 + shift_y, x1 + shift_x, y1 + shift_y
            ax.plot((x0, x1), (y0, y1), color=colors[kind],
                    lw=0.8 + 0.35 * math.sqrt(count),
                    alpha=0.42 if len(zones) == 1 else 0.68, zorder=2)
            ax.annotate("", xy=(x0 + 0.72 * dx, y0 + 0.72 * dy),
                        xytext=(x0 + 0.55 * dx, y0 + 0.55 * dy),
                        arrowprops={"arrowstyle": "-|>", "color": colors[kind],
                                    "lw": 1.25, "mutation_scale": 10}, zorder=3)
            if len(zones) > 1 and index <= len(zones):
                ax.text(x0 + 0.43 * dx, y0 + 0.43 * dy, str(index),
                        color=colors[kind], fontsize=7, weight="bold", zorder=4)
    for node, (x, y) in nodes.items():
        ax.scatter(x, y, s=62 if node == "O01" else 23,
                   color="#222222" if node == "O01" else "#565656", zorder=5)
        ax.annotate(node, (x, y), xytext=(3, 3),
                    textcoords="offset points", fontsize=7, zorder=6)
    ax.set(xlabel="投影东坐标（km）", ylabel="投影北坐标（km）")
    ax.set_aspect("equal", adjustable="datalim")
    ax.grid(alpha=0.2)
    handles = [Patch(facecolor=colors[kind], label=f"{kind}型") for kind in colors]
    ax.legend(handles=handles, loc="upper right", frameon=False, ncol=3)
    fig.tight_layout()
    for extension in ("png", "pdf"):
        fig.savefig(figure_dir / f"Q2_运输路线图.{extension}", dpi=220)
    plt.close(fig)


def save_tables(table_dir: Path, data_run: Path, sorties: list[dict[str, Any]],
                deliveries: list[dict[str, Any]], checks: list[dict[str, str]],
                construction: dict[str, Any]) -> dict[str, Any]:
    if table_dir.exists():
        raise FileExistsError(f"不覆盖已有结果：{table_dir}")
    validate_template()
    table_dir.mkdir(parents=True)
    sortie_rows = [{
        "架次编号": s["batch_id"], "无人机编号": s["uav_id"],
        "机型编号": s["type_id"], "电池编号": s["battery_id"],
        "开始时刻（s）": s["start_s"], "访问服务区顺序": s["zone_id"],
        "返回O01时刻（s）": s["return_s"], "架次能耗（kWh）": s["energy_kwh"],
    } for s in sorties]
    delivery_rows = [{
        "货箱编号": d["box_id"], "架次编号": d["batch_id"],
        "服务区编号": d["zone_id"], "交付完成时刻（s）": d["complete_s"],
    } for d in deliveries]
    write_csv(table_dir / "Q2_运输架次.csv", sortie_rows, SORTIE_COLUMNS)
    write_csv(table_dir / "Q2_逐箱交付.csv", delivery_rows, DELIVERY_COLUMNS)
    detailed_sorties = [{**s, "box_ids": ";".join(s["box_ids"])} for s in sorties]
    write_csv(table_dir / "Q2_架次与电池周转.csv", detailed_sorties, list(detailed_sorties[0]))
    write_csv(table_dir / "Q2_逐箱时限检查.csv", deliveries, list(deliveries[0]))
    write_csv(table_dir / "Q2_验收检查.csv", checks, ["check_id", "status", "detail"])
    failed = [check for check in checks if check["status"] == "FAIL"]
    hard_failed = [check for check in failed if check["check_id"].endswith(":hard_deadline")]
    hard_conflicts = sorted((d for d in deliveries if d["hard_slack_s"] is not None and
                             d["hard_slack_s"] < -TOL),
                            key=lambda d: (d["complete_s"], d["box_id"]))
    summary = {
        "status": "PASS" if not failed else "FAIL",
        "method": "same-zone ABC-capable FFD + deadline-first list scheduling",
        "baseline_method_id": "q2_baseline_abc_single_zone_v2",
        "baseline_not_optimized": True,
        "source_data_run": str(data_run),
        "source_manifest_sha256": sha256(data_run / "meta" / "validated_artifacts.json"),
        "hard_group_tasks": construction["hard_group_tasks"],
        "routine_tasks": construction["routine_tasks"],
        "type_sorties": dict(sorted(Counter(s["type_id"] for s in sorties).items())),
        "sorties": len(sorties), "delivered_boxes": len(deliveries),
        "uavs": 8, "batteries": 14,
        "makespan_s": max(s["return_s"] for s in sorties),
        "total_energy_kwh": sum(s["energy_kwh"] for s in sorties),
        "hard_deadline_boxes": sum(d["hard_deadline_s"] is not None for d in deliveries),
        "hard_deadline_violations": len(hard_failed),
        "first_failure": failed[0] if failed else None,
        "first_hard_conflict": ({"box_id": hard_conflicts[0]["box_id"],
                                  "delivered_s": hard_conflicts[0]["complete_s"],
                                  "deadline_s": hard_conflicts[0]["hard_deadline_s"]}
                                 if hard_conflicts else None),
        "soft_weighted_lateness_s": sum(
            d["priority"] * d["soft_lateness_s"] for d in deliveries
            if d["hard_deadline_s"] is None
        ),
        "battery_rule": "same-type shared pool; starts full; each used battery recharges to 100% before reuse",
        "charging_model": "two-stage 0-90% takes 65% of T_full; 90-100% takes 35%; parallel charging allowed",
        "note": "A failed hard deadline means this simple three-type baseline failed; it does not prove Q2 infeasible.",
    }
    (table_dir / "Q2_运行摘要.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    (table_dir / "Q2_运行说明.txt").write_text(
        "问题二基线：所有路线只访问一个服务区；硬时限箱同区组批，普通箱同区首次适应组批。\n"
        "每架次同时检查 A/B/C 的载荷、体积、能耗与返航余量，再按硬时限优先列表排程。\n"
        "三种机型均需实际使用；没有固定规定紧急箱只能由 A 型机运送。\n"
        "Q2_运输架次.csv 与 Q2_逐箱交付.csv 的列名与官方模板一致。\n"
        "逐箱交付时刻包含准备、装载、去程飞行、基础交接及本箱之前全部逐箱交接时间。\n"
        "电池返航后按官方两阶段规则充至 100% 才可复用；库存 6/4/4 已含初装。\n"
        "Q2_验收检查.csv 中如有 FAIL，本次基线不可行，但不能据此断定问题二无解。\n",
        encoding="utf-8")
    for filename, expected_columns, expected_count in (
        ("Q2_运输架次.csv", SORTIE_COLUMNS, len(sorties)),
        ("Q2_逐箱交付.csv", DELIVERY_COLUMNS, len(deliveries)),
    ):
        saved = pd.read_csv(table_dir / filename, encoding="utf-8-sig")
        if list(saved.columns) != expected_columns or len(saved) != expected_count:
            raise ValueError(f"官方格式表导出复核失败：{filename}")
    return summary


def save_figures(figure_dir: Path, sorties: list[dict[str, Any]],
                 deliveries: list[dict[str, Any]], uavs: list[UAV],
                 batteries: list[Battery], summary: dict[str, Any]) -> None:
    if figure_dir.exists():
        raise FileExistsError(f"不覆盖已有结果：{figure_dir}")
    os.environ.setdefault("MPLCONFIGDIR", str(CODE_DIR / ".mplcache"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    figure_dir.mkdir(parents=True)
    uav_ids = sorted(u.uav_id for u in uavs)
    battery_ids = sorted(b.battery_id for b in batteries)
    fig, axes = plt.subplots(2, 1, figsize=(12, 8), sharex=True,
                             gridspec_kw={"height_ratios": [len(uav_ids), len(battery_ids)]})
    kind_colors = {"A": "#3A80C1", "B": "#23A67A", "C": "#DA884A"}
    for s in sorties:
        color = kind_colors[s["type_id"]]
        axes[0].barh(uav_ids.index(s["uav_id"]), (s["return_s"] - s["start_s"]) / 3600,
                     left=s["start_s"] / 3600, height=0.65, color=color, edgecolor="white")
        y = battery_ids.index(s["battery_id"])
        axes[1].barh(y, (s["return_s"] - s["start_s"]) / 3600,
                     left=s["start_s"] / 3600, height=0.65, color=color, edgecolor="white")
        axes[1].barh(y, s["charge_s"] / 3600, left=s["return_s"] / 3600,
                     height=0.65, color="#AEB5BD", edgecolor="white")
    for ax, labels, title in ((axes[0], uav_ids, "Transport UAVs"),
                              (axes[1], battery_ids, "Shared batteries: use and recharge")):
        ax.set_yticks(range(len(labels)), labels)
        ax.invert_yaxis()
        ax.set_title(title, loc="left", fontsize=11)
        ax.grid(axis="x", alpha=0.25)
        ax.set_axisbelow(True)
    axes[1].set_xlabel("Time since start (h)")
    fig.legend(handles=[Patch(facecolor=color, label=f"Type {kind}")
                        for kind, color in kind_colors.items()] +
               [Patch(facecolor="#AEB5BD", label="Charging")],
               loc="lower center", ncol=4, frameon=False)
    fig.subplots_adjust(bottom=0.10, hspace=0.32, left=0.14, right=0.98, top=0.94)
    for extension in ("png", "pdf"):
        fig.savefig(figure_dir / f"Q2_机电池时间线.{extension}", dpi=220)
    plt.close(fig)

    hard = sorted((d for d in deliveries if d["hard_deadline_s"] is not None),
                  key=lambda d: (d["hard_slack_s"], d["box_id"]))
    fig, ax = plt.subplots(figsize=(11, 5.5))
    ax.bar(range(len(hard)), [d["hard_slack_s"] / 60 for d in hard],
           color=["#C84D4D" if d["hard_slack_s"] < -TOL else "#3A80C1" for d in hard])
    ax.axhline(0, color="#333333", linewidth=0.9)
    ax.set_xlabel("Hard-deadline boxes ordered by slack")
    ax.set_ylabel("Deadline slack (min)")
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    fig.tight_layout()
    for extension in ("png", "pdf"):
        fig.savefig(figure_dir / f"Q2_硬时限余量.{extension}", dpi=220)
    plt.close(fig)
    save_delivery_figures(figure_dir, deliveries, plt)
    save_route_figure(figure_dir, Path(summary["source_data_run"]), sorties, plt)
    (figure_dir / "图表说明.txt").write_text(
        "Q2_机电池时间线：上图为 8 架实体机占用，下图为共享电池占用与充电。\n"
        "Q2_硬时限余量：每个有硬截止的箱子以截止时刻减实际交付时刻作图；红色为超时。\n"
        "Q2_逐箱交付进度：累计交接完成箱数；虚线为硬截止时刻。\n"
        "Q2_交付硬时限对照：每箱以交接完成时刻与硬截止比较；点在对角线下方为按时。\n"
        "Q2_运输路线图：投影坐标下的飞行方向，重复路线合并；多区路线数字为访问顺序。\n"
        f"本次基线状态：{summary['status']}；如为 FAIL，图仅用于定位冲突，不能作为可行方案。\n",
        encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="D 题问题二：A/B/C 同区组批与共享电池列表排程基线")
    parser.add_argument("--data-run", type=Path, help="已验收的数据运行目录；默认最新一次")
    parser.add_argument("--output-root", type=Path, help="结果根目录；默认 code/2_outputs")
    args = parser.parse_args()
    data_run = args.data_run.resolve() if args.data_run else find_latest_validated_run()
    if not verify_ready(data_run):
        raise ValueError(f"数据目录未通过验收：{data_run}")
    tasks, boxes_by_id, construction = make_tasks(data_run)
    uavs, batteries = read_resources(data_run)
    sorties, deliveries = schedule_baseline(tasks, uavs, batteries)
    checks = check_plan(sorties, deliveries, tasks, uavs, batteries, boxes_by_id)
    if not verify_ready(data_run):
        raise ValueError("排程期间输入数据发生改变")
    output_root = args.output_root.resolve() if args.output_root else CODE_DIR / "2_outputs"
    stamp = datetime.now().strftime("%y%m%d_%H%M%S")
    method_dir = output_root / "0_baseline"
    table_dir = method_dir / f"q2_base_{stamp}_table"
    figure_dir = method_dir / f"q2_base_{stamp}_figure"
    if table_dir.exists() or figure_dir.exists():
        raise FileExistsError("本次结果目录已存在；请稍后重试或指定新的 --output-root")
    summary = save_tables(table_dir, data_run, sorties, deliveries, checks, construction)
    save_figures(figure_dir, sorties, deliveries, uavs, batteries, summary)
    (table_dir / "Q2_基线状态.txt").write_text(
        f"{summary['status']}\nfirst_failure={summary['first_failure']}\n"
        f"first_hard_conflict={summary['first_hard_conflict']}\n", encoding="utf-8")
    print(f"Q2 baseline {summary['status']}: {summary['sorties']} sorties, "
          f"{summary['delivered_boxes']} boxes, {summary['hard_deadline_violations']} hard-deadline violations")
    print(f"makespan={summary['makespan_s']:.3f} s, total_energy={summary['total_energy_kwh']:.6f} kWh")
    if summary["first_failure"]:
        print(f"first failure: {summary['first_failure']}")
    if summary["first_hard_conflict"]:
        print(f"first hard conflict: {summary['first_hard_conflict']}")
    print(f"tables: {table_dir}")
    print(f"figures: {figure_dir}")


if __name__ == "__main__":
    main()
