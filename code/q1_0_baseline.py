"""D 题第一问可行基线：安全载荷、单箱保底、同区 FFD 与回代验收。

注意：水平与爬升能耗的具体分项公式来自组内执行思路，是建模假设，
而非官方题面直接给出的公式。这里不求解集合划分优化，也不排实体机/电池。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import numpy as np
import pytest
from openpyxl import load_workbook

os.environ.setdefault("MPLCONFIGDIR", str(Path(__file__).resolve().parent / ".mplcache"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import colors, font_manager

from _0_pipeline import sha256, verify_ready


CODE_DIR = Path(__file__).resolve().parent
PROJECT = CODE_DIR.parent
OUTPUTS = CODE_DIR / "1_outputs"
FIGURES = OUTPUTS
GRAVITY = 9.81
KWH_JOULES = 3_600_000.0
TOL = 1e-9
TEMPLATE_COLUMNS = [
    "架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）",
    "总体积（m³）", "往返时间（s）", "架次能耗（kWh）", "返航SOC（%）",
]


def equivalent_range_m(model: dict[str, Any], load_kg: float) -> float:
    capacity = float(model["max_payload_kg"])
    if capacity <= 0 or load_kg < -TOL or load_kg > capacity + TOL:
        raise ValueError(f"非法载荷：{load_kg} kg；额定载重 {capacity} kg")
    q = max(0.0, min(float(load_kg), capacity))
    result = float(model["range_empty_m"]) - (
        float(model["range_empty_m"]) - float(model["range_full_m"])
    ) * (q / capacity) ** 1.5
    if result <= 0:
        raise ValueError("等效航程必须为正")
    return result


def leg_energy_kwh(model: dict[str, Any], arc: dict[str, Any], load_kg: float) -> dict[str, float]:
    horizontal = (float(model["energy_use_kwh"]) * float(arc["distance_m"]) /
                  equivalent_range_m(model, load_kg))
    climb = ((float(model["empty_mass_kg"]) + load_kg) * GRAVITY *
             float(arc["climb_m"]) /
             (KWH_JOULES * float(model["climb_efficiency"])))
    return {"horizontal_kwh": horizontal, "climb_kwh": climb,
            "total_kwh": horizontal + climb}


def leg_time_s(model: dict[str, Any], arc: dict[str, Any]) -> float:
    return (float(arc["climb_m"]) / float(model["climb_speed_mps"]) +
            float(arc["distance_m"]) / float(model["cruise_speed_mps"]) +
            float(arc["descent_m"]) / float(model["descent_speed_mps"]))


def roundtrip_energy_kwh(model: dict[str, Any], outward: dict[str, Any],
                         homeward: dict[str, Any], load_kg: float) -> float:
    # 返程货物已交付，载荷始终为 0；使用反向航段自身的爬升高度。
    return (leg_energy_kwh(model, outward, load_kg)["total_kwh"] +
            leg_energy_kwh(model, homeward, 0.0)["total_kwh"])


def max_safe_payload_kg(model: dict[str, Any], outward: dict[str, Any],
                        homeward: dict[str, Any]) -> dict[str, Any]:
    rho = float(model["return_soc_min"])
    if not 0 <= rho < 1:
        raise ValueError(f"返航余量比例超范围：{rho}")
    rated = float(model["max_payload_kg"])
    limit = (1.0 - rho) * float(model["energy_use_kwh"])
    empty_energy = roundtrip_energy_kwh(model, outward, homeward, 0.0)
    result: dict[str, Any] = {
        "energy_limit_kwh": limit, "energy_empty_kwh": empty_energy,
        "rated_payload_kg": rated,
    }
    if empty_energy > limit + TOL:
        return dict(result, status="unreachable", q_max_kg=None,
                    energy_at_qmax_kwh=None, soc_at_qmax_pct=None)
    rated_energy = roundtrip_energy_kwh(model, outward, homeward, rated)
    if rated_energy <= limit + TOL:
        safe_load, status = rated, "rated_payload"
    else:
        low, high = 0.0, rated
        for _ in range(80):
            mid = (low + high) / 2.0
            if roundtrip_energy_kwh(model, outward, homeward, mid) <= limit:
                low = mid
            else:
                high = mid
        # 对外保留 6 位小数并向下截断，避免打印舍入越过能量边界。
        safe_load, status = math.floor(low * 1_000_000) / 1_000_000, "energy_limited"
    energy = roundtrip_energy_kwh(model, outward, homeward, safe_load)
    if energy > limit + TOL:
        raise AssertionError("二分的安全载荷未通过能量回代")
    return dict(result, status=status, q_max_kg=safe_load,
                energy_at_qmax_kwh=energy,
                soc_at_qmax_pct=100.0 * (1.0 - energy / float(model["energy_use_kwh"])))


def sortie_metrics(model: dict[str, Any], outward: dict[str, Any],
                   homeward: dict[str, Any], boxes: list[dict[str, Any]]) -> dict[str, Any]:
    if not boxes:
        raise ValueError("架次不能没有货箱")
    mass = sum(float(box["mass_kg"]) for box in boxes)
    volume = sum(float(box["volume_m3"]) for box in boxes)
    outbound = leg_energy_kwh(model, outward, mass)
    inbound = leg_energy_kwh(model, homeward, 0.0)
    energy = outbound["total_kwh"] + inbound["total_kwh"]
    energy_use = float(model["energy_use_kwh"])
    energy_limit = (1.0 - float(model["return_soc_min"])) * energy_use
    flight_out = leg_time_s(model, outward)
    flight_back = leg_time_s(model, homeward)
    prep = float(model["prep_s"])
    loading = len(boxes) * float(model["load_per_box_s"])
    handoff = (float(model["handoff_base_s"]) +
               len(boxes) * float(model["handoff_per_box_s"]))
    failures = []
    if mass > float(model["max_payload_kg"]) + TOL:
        failures.append("rated_mass")
    if volume > float(model["capacity_m3"]) + TOL:
        failures.append("volume")
    if energy > energy_limit + TOL:
        failures.append("return_energy")
    return {
        "feasible": not failures, "failure_reasons": ";".join(failures),
        "mass_kg": mass, "volume_m3": volume, "box_count": len(boxes),
        "energy_out_kwh": outbound["total_kwh"], "energy_back_kwh": inbound["total_kwh"],
        "energy_kwh": energy, "energy_limit_kwh": energy_limit,
        "return_soc": 1.0 - energy / energy_use,
        "flight_out_s": flight_out, "flight_back_s": flight_back,
        "flight_s": flight_out + flight_back,
        "prep_s": prep, "loading_s": loading, "handoff_s": handoff,
        "operation_s": prep + loading + flight_out + handoff + flight_back,
    }


def best_type_for_batch(boxes: list[dict[str, Any]], models: dict[str, dict[str, Any]],
                        outward: dict[str, Any], homeward: dict[str, Any]) -> dict[str, Any] | None:
    best = None
    for type_id in sorted(models):
        model = models[type_id]
        mass = sum(float(box["mass_kg"]) for box in boxes)
        volume = sum(float(box["volume_m3"]) for box in boxes)
        if (mass > float(model["max_payload_kg"]) + TOL or
                volume > float(model["capacity_m3"]) + TOL):
            continue
        result = sortie_metrics(model, outward, homeward, boxes)
        if not result["feasible"]:
            continue
        candidate = dict(result, type_id=type_id)
        if best is None or candidate["energy_kwh"] < best["energy_kwh"] - 1e-12:
            best = candidate
        # 已按型号排序；能耗平局时保留先遇到的较小型号。
    return best


def _read_inputs(data_run: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]],
                                            dict[tuple[str, str], dict[str, Any]]]:
    if not verify_ready(data_run):
        raise ValueError(f"输入数据未通过哈希验收，不可建模：{data_run}")
    clean = data_run / "clean"
    boxes_frame = pd.read_csv(clean / "boxes.csv", dtype={"box_id": str, "zone_id": str})
    types_frame = pd.read_csv(clean / "transport_types.csv", dtype={"type_id": str})
    arcs_frame = pd.read_csv(data_run / "derived" / "arc_geometry.csv",
                             dtype={"from_id": str, "to_id": str})
    if len(boxes_frame) != 80 or len(types_frame) != 3 or len(arcs_frame) != 240:
        raise ValueError("输入表规模与正式附件不符")
    boxes = boxes_frame.to_dict("records")
    models = {str(row["type_id"]): row for row in types_frame.to_dict("records")}
    arcs = {(str(row["from_id"]), str(row["to_id"])): row
            for row in arcs_frame.to_dict("records")}
    return boxes, models, arcs


def _plan_from_groups(prefix: str, groups: list[dict[str, Any]],
                      models: dict[str, dict[str, Any]],
                      arcs: dict[tuple[str, str], dict[str, Any]]) -> list[dict[str, Any]]:
    plan = []
    for sequence, group in enumerate(groups, 1):
        zone = group["zone_id"]
        batch_boxes = group["boxes"]
        chosen = best_type_for_batch(batch_boxes, models, arcs[("O01", zone)], arcs[(zone, "O01")])
        if chosen is None:
            raise ValueError(f"{zone} 的货箱组不可由任一机型安全往返：{[b['box_id'] for b in batch_boxes]}")
        plan.append(dict(chosen, batch_id=f"{prefix}{sequence:03d}", zone_id=zone,
                         box_ids=[str(box["box_id"]) for box in batch_boxes]))
    return plan


def _validate_plan(label: str, plan: list[dict[str, Any]],
                   source_boxes: list[dict[str, Any]], models: dict[str, dict[str, Any]],
                   arcs: dict[tuple[str, str], dict[str, Any]]) -> list[dict[str, str]]:
    checks: list[dict[str, str]] = []
    source = {str(box["box_id"]): box for box in source_boxes}
    seen = Counter(box_id for batch in plan for box_id in batch["box_ids"])

    def check(name: str, good: bool, detail: str = "") -> None:
        checks.append({"check_id": f"{label}:{name}",
                       "status": "PASS" if good else "FAIL", "detail": detail})

    check("all_boxes_exactly_once", seen == Counter(source.keys()),
          f"assigned={sum(seen.values())}, unique={len(seen)}, expected={len(source)}")
    for batch in plan:
        batch_id, zone, type_id = batch["batch_id"], batch["zone_id"], batch["type_id"]
        if any(box_id not in source for box_id in batch["box_ids"]):
            check(f"{batch_id}:known_box_ids", False)
            continue
        group = [source[box_id] for box_id in batch["box_ids"]]
        check(f"{batch_id}:one_zone", all(box["zone_id"] == zone for box in group))
        recalculated = sortie_metrics(models[type_id], arcs[("O01", zone)],
                                      arcs[(zone, "O01")], group)
        check(f"{batch_id}:mass_volume_energy", recalculated["feasible"],
              recalculated["failure_reasons"])
        fields = ("mass_kg", "volume_m3", "energy_kwh", "return_soc", "flight_s", "operation_s")
        check(f"{batch_id}:replay", all(math.isclose(batch[field], recalculated[field],
                                                     rel_tol=0, abs_tol=1e-8) for field in fields))
    return checks


def _metrics(label: str, plan: list[dict[str, Any]]) -> dict[str, Any]:
    return {"baseline": label, "sorties": len(plan),
            "delivered_boxes": sum(len(batch["box_ids"]) for batch in plan),
            "total_mass_kg": sum(batch["mass_kg"] for batch in plan),
            "total_volume_m3": sum(batch["volume_m3"] for batch in plan),
            "total_energy_kwh": sum(batch["energy_kwh"] for batch in plan),
            "cumulative_operation_s": sum(batch["operation_s"] for batch in plan),
            "cumulative_flight_s": sum(batch["flight_s"] for batch in plan),
            "min_return_soc_pct": min(batch["return_soc"] for batch in plan) * 100}


def run_model(data_run: Path) -> dict[str, Any]:
    data_run = data_run.resolve()
    boxes, models, arcs = _read_inputs(data_run)
    zones = sorted({str(box["zone_id"]) for box in boxes})
    if len(zones) != 15 or set(models) != {"A", "B", "C"}:
        raise ValueError("第一问要求 15 个服务区与 A/B/C 三种机型")
    for zone in zones:
        for leg in (("O01", zone), (zone, "O01")):
            if leg not in arcs or not bool(arcs[leg]["valid_flag"]):
                raise ValueError(f"第一问必要航段无有效地形数据：{leg}")
    max_payloads = []
    checks: list[dict[str, str]] = []
    for type_id in sorted(models):
        model = models[type_id]
        for zone in zones:
            outward, homeward = arcs[("O01", zone)], arcs[(zone, "O01")]
            result = max_safe_payload_kg(model, outward, homeward)
            max_payloads.append(dict(type_id=type_id, zone_id=zone,
                                     out_distance_m=outward["distance_m"],
                                     back_distance_m=homeward["distance_m"], **result))
            if result["status"] == "unreachable":
                good = (result["q_max_kg"] is None and
                        result["energy_empty_kwh"] > result["energy_limit_kwh"] + TOL)
            else:
                good = result["energy_at_qmax_kwh"] <= result["energy_limit_kwh"] + TOL
            checks.append({"check_id": f"max_payload:{type_id}:{zone}",
                           "status": "PASS" if good else "FAIL",
                           "detail": result["status"]})
    checks.append({"check_id": "max_payloads:45_rows", "status": "PASS" if len(max_payloads) == 45 else "FAIL",
                   "detail": str(len(max_payloads))})

    single_groups = []
    ffd_groups = []
    for zone in zones:
        local = sorted((box for box in boxes if box["zone_id"] == zone), key=lambda b: str(b["box_id"]))
        for box in local:
            if best_type_for_batch([box], models, arcs[("O01", zone)], arcs[(zone, "O01")]) is None:
                raise ValueError(f"单箱无可行机型：{box['box_id']}")
            single_groups.append({"zone_id": zone, "boxes": [box]})
        ordered = sorted(local, key=lambda b: (-float(b["mass_kg"]), -float(b["volume_m3"]), str(b["box_id"])))
        local_groups: list[dict[str, Any]] = []
        for box in ordered:
            placed = False
            for group in local_groups:
                candidate = group["boxes"] + [box]
                if best_type_for_batch(candidate, models, arcs[("O01", zone)], arcs[(zone, "O01")]) is not None:
                    group["boxes"].append(box)
                    placed = True
                    break
            if not placed:
                local_groups.append({"zone_id": zone, "boxes": [box]})
        ffd_groups.extend(local_groups)

    single = _plan_from_groups("S", single_groups, models, arcs)
    ffd = _plan_from_groups("F", ffd_groups, models, arcs)
    checks.extend(_validate_plan("single_box", single, boxes, models, arcs))
    checks.extend(_validate_plan("ffd", ffd, boxes, models, arcs))
    checks.append({"check_id": "ffd:sorties_not_above_single_box",
                   "status": "PASS" if len(ffd) <= len(single) else "FAIL",
                   "detail": f"ffd={len(ffd)}, single={len(single)}"})
    metrics = [_metrics("single_box", single), _metrics("ffd", ffd)]
    for metric in metrics:
        checks.append({"check_id": f"{metric['baseline']}:reconcile_totals",
                       "status": "PASS" if (metric["delivered_boxes"] == 80 and
                           math.isclose(metric["total_mass_kg"], 758, abs_tol=1e-8) and
                           math.isclose(metric["total_volume_m3"], 2.011, abs_tol=1e-8)) else "FAIL",
                       "detail": str(metric["delivered_boxes"])})
    return {"data_run": data_run, "max_payloads": max_payloads,
            "single_box": single, "ffd": ffd, "metrics": metrics, "checks": checks}


def _template_rows(plan: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{
        "架次编号": batch["batch_id"], "服务区编号": batch["zone_id"],
        "机型编号": batch["type_id"], "货箱编号列表": ";".join(batch["box_ids"]),
        "总质量（kg）": batch["mass_kg"], "总体积（m³）": batch["volume_m3"],
        "往返时间（s）": batch["operation_s"], "架次能耗（kWh）": batch["energy_kwh"],
        "返航SOC（%）": batch["return_soc"] * 100,
    } for batch in plan]


def _detail_rows(plan: list[dict[str, Any]]) -> list[dict[str, Any]]:
    keys = ("batch_id", "zone_id", "type_id", "box_ids", "box_count", "mass_kg",
            "volume_m3", "energy_out_kwh", "energy_back_kwh", "energy_kwh",
            "energy_limit_kwh", "return_soc", "flight_out_s", "flight_back_s",
            "flight_s", "prep_s", "loading_s", "handoff_s", "operation_s")
    return [{key: ";".join(batch[key]) if key == "box_ids" else batch[key] for key in keys}
            for batch in plan]


def _save_csv(path: Path, rows: list[dict[str, Any]], columns: list[str] | None = None) -> None:
    pd.DataFrame(rows, columns=columns).to_csv(path, index=False, encoding="utf-8-sig", float_format="%.12g")


def _verify_written_outputs(result: dict[str, Any], output_dir: Path) -> None:
    """复读落盘结果；写出过程有错时不得留下 Q1_READY 标记。"""
    for filename, plan in (("Q1_单箱保底.csv", result["single_box"]),
                           ("Q1_单点组批.csv", result["ffd"])):
        frame = pd.read_csv(output_dir / filename, dtype={column: str for column in
                            TEMPLATE_COLUMNS[:4]}, encoding="utf-8-sig")
        expected = _template_rows(plan)
        if list(frame.columns) != TEMPLATE_COLUMNS or len(frame) != len(expected):
            raise ValueError(f"导出复核失败：{filename} 的列名或行数不符")
        for index, row in enumerate(expected):
            actual = frame.iloc[index]
            for column in TEMPLATE_COLUMNS[:4]:
                if actual[column] != row[column]:
                    raise ValueError(f"导出复核失败：{filename} 第 {index + 2} 行 {column} 不符")
            for column in TEMPLATE_COLUMNS[4:]:
                if not math.isclose(float(actual[column]), float(row[column]),
                                    rel_tol=1e-11, abs_tol=1e-8):
                    raise ValueError(f"导出复核失败：{filename} 第 {index + 2} 行 {column} 不符")

    payloads = pd.read_csv(output_dir / "Q1_安全载荷.csv", encoding="utf-8-sig")
    expected_pairs = {(row["type_id"], row["zone_id"]) for row in result["max_payloads"]}
    actual_pairs = set(zip(payloads["type_id"], payloads["zone_id"]))
    if len(payloads) != 45 or actual_pairs != expected_pairs:
        raise ValueError("导出复核失败：安全载荷矩阵不是预期的 45 个机型-服务区组合")
    for label, filename in (("指标", "Q1_指标对照.csv"), ("验收", "Q1_验收检查.csv")):
        frame = pd.read_csv(output_dir / filename, encoding="utf-8-sig")
        expected_count = len(result["metrics"] if label == "指标" else result["checks"])
        if len(frame) != expected_count:
            raise ValueError(f"导出复核失败：{label}行数不符")
        if label == "验收" and not frame["status"].eq("PASS").all():
            raise ValueError("导出复核失败：验收文件存在失败项")


def save_results(result: dict[str, Any], output_dir: Path) -> None:
    if output_dir.exists():
        raise FileExistsError(f"不覆盖既有结果：{output_dir}")
    if any(check["status"] != "PASS" for check in result["checks"]):
        raise ValueError("第一问回代检查未全通过，结果不发布")
    if not verify_ready(result["data_run"]):
        raise ValueError("输入数据在求解期间发生改变")
    template = load_workbook(PROJECT / "0_3_数据" / "结果提交模板.xlsx", read_only=True, data_only=True)
    actual_header = [template["Q1_单点组批"].cell(1, col).value for col in range(1, 10)]
    template.close()
    if actual_header != TEMPLATE_COLUMNS:
        raise ValueError("官方 Q1 模板列名变化，停止导出以免错列")
    output_dir.mkdir(parents=True)
    _save_csv(output_dir / "Q1_安全载荷.csv", result["max_payloads"])
    _save_csv(output_dir / "Q1_单箱保底.csv", _template_rows(result["single_box"]), TEMPLATE_COLUMNS)
    _save_csv(output_dir / "Q1_单点组批.csv", _template_rows(result["ffd"]), TEMPLATE_COLUMNS)
    _save_csv(output_dir / "Q1_批次明细.csv", _detail_rows(result["ffd"]))
    _save_csv(output_dir / "Q1_指标对照.csv", result["metrics"])
    _save_csv(output_dir / "Q1_验收检查.csv", result["checks"])
    assumptions = {
        "source_data_run": str(result["data_run"]),
        "source_validated_artifacts_sha256": sha256(result["data_run"] / "meta" / "validated_artifacts.json"),
        "method": "single-box fallback and same-zone first-fit-decreasing; not optimized",
        "energy_formula_status": "modeling assumption from execution guide page 6, not an explicit official formula",
        "energy_horizontal": "E_use * distance_m / [L0 - (L0-LF)*(q/Q)^1.5]",
        "energy_climb": "(M_empty+q)*9.81*climb_m/(3600000*climb_efficiency)",
        "return_safety": "sortie_energy <= (1-rho)*E_use; return_SOC = 1-sortie_energy/E_use",
        "route": "O01 -> one service zone -> O01; outbound actual payload; return zero payload",
        "time_in_template": "full operation time = prep + per-box loading + both flight legs + one base handoff + per-box handoff",
        "cumulative_operation_time": "sum of full operation time over all sorties; not makespan",
        "excluded": ["entity-UAV scheduling", "battery scheduling", "delivery deadlines",
                     "exact set partitioning optimization", "return-margin sensitivity"],
        "metrics": result["metrics"],
        "checks_total": len(result["checks"]),
        "checks_passed": sum(check["status"] == "PASS" for check in result["checks"]),
    }
    (output_dir / "Q1_运行摘要.json").write_text(json.dumps(assumptions, ensure_ascii=False, indent=2),
                                                  encoding="utf-8")
    explanation = (
        "D题第一问：可行基线，不是最优解。\n"
        "按官方题面只做单服务区 O01→Si→O01，不排实体无人机、共享电池和时限。\n"
        "Q1_安全载荷.csv：A/B/C×15区的45个安全载荷边界。\n"
        "Q1_单箱保底.csv：80箱各占一架次的保底方案。\n"
        "Q1_单点组批.csv：FFD组批，列名与官方Q1模板A:I完全一致。\n"
        "Q1_批次明细.csv：把飞行、准备、装载、交接、能耗分项列开，便于复算。\n"
        "模板的“往返时间”采用完整架次作业时间；纯飞行时间见批次明细。\n"
        "水平/爬升具体能耗公式是执行思路的建模假设，不是官方直接给出的公式。\n"
        "不包含集合划分优化、返航余量敏感性；这两项是第一问后续工作。\n"
    )
    (output_dir / "Q1_运行说明.txt").write_text(explanation, encoding="utf-8")
    _verify_written_outputs(result, output_dir)
    (output_dir / "Q1_READY.txt").write_text(
        f"PASS; {len(result['checks'])} checks; source={result['data_run']}\n", encoding="utf-8")


def find_latest_validated_run() -> Path:
    for path in sorted((CODE_DIR / "0_outputs").glob("run_*"), reverse=True):
        if path.is_dir() and verify_ready(path):
            return path
    raise FileNotFoundError("找不到通过验收的数据目录，请先运行 _0_pipeline.py")


def figure_dir_for(table_dir: Path) -> Path:
    """让同一次运行的表格、图片使用相同时间戳与原有目录格式。"""
    name = table_dir.name
    return table_dir.with_name(name[:-6] + "_figure" if name.endswith("_table") else name + "_figure")


SOURCE_FILES = ["Q1_安全载荷.csv", "Q1_单箱保底.csv", "Q1_单点组批.csv",
                "Q1_指标对照.csv", "Q1_验收检查.csv", "Q1_运行摘要.json", "Q1_READY.txt"]
GRAY = "#526174"
TEAL = "#007D82"
GRID = "#D8DEE5"


def latest_result_dir() -> Path:
    for path in sorted(OUTPUTS.glob("q1_base_*_table"), reverse=True):
        if path.is_dir() and (path / "Q1_READY.txt").is_file():
            return path
    raise FileNotFoundError("找不到有 Q1_READY.txt 的第一问结果；请先运行 q1_0_baseline.py")


def _read(path: Path) -> pd.DataFrame:
    if not path.is_file():
        raise FileNotFoundError(f"缺少图表数据：{path}")
    return pd.read_csv(path, encoding="utf-8-sig")


def _box_ids(frame: pd.DataFrame) -> list[str]:
    return [box_id for cell in frame["货箱编号列表"].astype(str)
            for box_id in cell.split(";") if box_id]


def load_results(result_dir: Path) -> dict[str, object]:
    result_dir = result_dir.resolve()
    if not (result_dir / "Q1_READY.txt").is_file():
        raise ValueError("结果缺少 Q1_READY.txt，不应作图")
    payload = _read(result_dir / "Q1_安全载荷.csv")
    single = _read(result_dir / "Q1_单箱保底.csv")
    ffd = _read(result_dir / "Q1_单点组批.csv")
    metrics = _read(result_dir / "Q1_指标对照.csv")
    qa = _read(result_dir / "Q1_验收检查.csv")
    for name, frame in (("单箱保底", single), ("同区FFD", ffd)):
        if list(frame.columns) != TEMPLATE_COLUMNS:
            raise ValueError(f"{name} 的列名或顺序与第一问模板不一致")
    if qa.empty or not qa["status"].eq("PASS").all():
        raise ValueError("第一问验收表中存在失败项，不应作图")
    if set(metrics["baseline"]) != {"single_box", "ffd"} or len(metrics) != 2:
        raise ValueError("指标对照表应恰有两种基线")
    metric_by_name = metrics.set_index("baseline")
    for name, frame in (("single_box", single), ("ffd", ffd)):
        metric = metric_by_name.loc[name]
        if len(frame) != int(metric["sorties"]):
            raise ValueError(f"{name} 的架次行数与指标不一致")
        for column, expected in (("总质量（kg）", "total_mass_kg"),
                                 ("总体积（m³）", "total_volume_m3"),
                                 ("往返时间（s）", "cumulative_operation_s"),
                                 ("架次能耗（kWh）", "total_energy_kwh")):
            if not math.isclose(float(frame[column].sum()), float(metric[expected]),
                                rel_tol=1e-9, abs_tol=1e-5):
                raise ValueError(f"{name} 的 {column} 与指标对照表不一致")
        ids = _box_ids(frame)
        if len(ids) != 80 or len(set(ids)) != 80:
            raise ValueError(f"{name} 未将 80 箱恰好安排一次")
    if set(_box_ids(single)) != set(_box_ids(ffd)):
        raise ValueError("两个基线的货箱集合不相同")
    expected_pairs = {(kind, f"S{zone:03d}") for kind in "ABC" for zone in range(1, 16)}
    actual_pairs = set(zip(payload["type_id"], payload["zone_id"]))
    if len(payload) != 45 or actual_pairs != expected_pairs:
        raise ValueError("安全载荷表不包含完整的 3×15 组合")
    allowed = {"rated_payload", "energy_limited", "unreachable"}
    if not set(payload["status"]).issubset(allowed):
        raise ValueError("安全载荷状态存在未知值")
    reachable = payload["status"] != "unreachable"
    if (payload.loc[reachable, "q_max_kg"].isna().any() or
            (payload.loc[reachable, "q_max_kg"] < 0).any() or
            (payload.loc[reachable, "q_max_kg"] >
             payload.loc[reachable, "rated_payload_kg"] + 1e-8).any() or
            payload.loc[~reachable, "q_max_kg"].notna().any()):
        raise ValueError("安全载荷数值与状态不一致")
    summary = json.loads((result_dir / "Q1_运行摘要.json").read_text(encoding="utf-8"))
    if int(summary["checks_passed"]) != len(qa):
        raise ValueError("运行摘要与验收表不一致")
    return {"directory": result_dir, "payload": payload, "single": single,
            "ffd": ffd, "metrics": metric_by_name, "qa_count": len(qa),
            "summary": summary}


def configure_plot_style() -> str:
    candidates = [Path("C:/Windows/Fonts/simhei.ttf"),
                  Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")]
    for font_path in candidates:
        if font_path.is_file():
            font_manager.fontManager.addfont(str(font_path))
            font_name = font_manager.FontProperties(fname=str(font_path)).get_name()
            break
    else:
        try:
            font_name = font_manager.FontProperties(
                fname=font_manager.findfont("Noto Sans CJK SC", fallback_to_default=False)
            ).get_name()
        except ValueError as error:
            raise RuntimeError("找不到中文字体，无法安全输出中文图表") from error
    plt.rcParams.update({
        "font.family": font_name, "font.size": 10, "axes.unicode_minus": False,
        "pdf.fonttype": 42, "ps.fonttype": 42, "figure.facecolor": "white",
        "axes.facecolor": "white", "savefig.facecolor": "white",
    })
    return font_name


def _save_figure(fig: plt.Figure, destination: Path) -> None:
    fig.savefig(destination.with_suffix(".pdf"), dpi=300)
    fig.savefig(destination.with_suffix(".png"), dpi=300)
    plt.close(fig)


def plot_comparison(data: dict[str, object], destination: Path) -> dict[str, float]:
    metrics = data["metrics"]
    single, ffd = metrics.loc["single_box"], metrics.loc["ffd"]
    definitions = [
        ("架次数", "架次", float(single["sorties"]), float(ffd["sorties"]), "{:.0f}"),
        ("总运输能耗", "kWh", float(single["total_energy_kwh"]),
         float(ffd["total_energy_kwh"]), "{:.1f}"),
        ("累计作业时间", "小时", float(single["cumulative_operation_s"]) / 3600,
         float(ffd["cumulative_operation_s"]) / 3600, "{:.2f}"),
    ]
    fig = plt.figure(figsize=(10.8, 4.2), layout="constrained")
    grid = fig.add_gridspec(2, 3, height_ratios=[14, 1.2])
    axes = [fig.add_subplot(grid[0, column]) for column in range(3)]
    note = fig.add_subplot(grid[1, :])
    note.axis("off")
    reductions: dict[str, float] = {}
    for ax, (label, unit, initial, batched, fmt) in zip(axes, definitions):
        reduction = 100 * (1 - batched / initial)
        reductions[label] = reduction
        bars = ax.bar([0, 1], [initial, batched], width=0.56,
                      color=[GRAY, TEAL], edgecolor="#283844", linewidth=0.8,
                      hatch=["///", ""])
        ax.set_xticks([0, 1], ["单箱保底", "同区FFD"])
        ax.set_ylabel(f"{label}（{unit}）")
        ax.set_ylim(0, max(initial, batched) * 1.32)
        ax.grid(axis="y", color=GRID, linewidth=0.65)
        ax.set_axisbelow(True)
        ax.spines[["top", "right"]].set_visible(False)
        for bar, value in zip(bars, (initial, batched)):
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() +
                    max(initial, batched) * 0.025, fmt.format(value),
                    ha="center", va="bottom", fontsize=10, fontweight="bold")
        ax.text(0.98, 0.97, f"减少 {reduction:.1f}%", transform=ax.transAxes,
                ha="right", va="top", color=TEAL, fontsize=9)
    note.text(0, 0.45,
              "注：两方案均为可行基线，非最优解；累计作业时间为各架次时间之和，不是任务完工时刻。",
              ha="left", va="center", fontsize=9, color="#364451")
    _save_figure(fig, destination)
    return reductions


def plot_safe_payload(data: dict[str, object], destination: Path) -> dict[str, object]:
    source = data["payload"].set_index(["type_id", "zone_id"])
    zones = [f"S{i:03d}" for i in range(1, 16)]
    values = np.full((3, 15), np.nan, dtype=float)
    labels: list[list[str]] = []
    limited: list[str] = []
    for row, kind in enumerate("ABC"):
        row_labels = []
        for column, zone in enumerate(zones):
            item = source.loc[(kind, zone)]
            if item["status"] == "unreachable":
                row_labels.append("不可达")
                continue
            qmax, rated = float(item["q_max_kg"]), float(item["rated_payload_kg"])
            values[row, column] = qmax / rated
            if item["status"] == "energy_limited":
                limited.append(f"{kind}-{zone}")
            shown = f"{qmax:.1f}" if not math.isclose(qmax, rated, abs_tol=1e-6) else f"{rated:.0f}"
            row_labels.append(shown + ("*" if item["status"] == "energy_limited" else ""))
        labels.append(row_labels)
    cmap = matplotlib.colormaps["YlGnBu"].copy()
    cmap.set_bad("#E6E9ED")
    fig = plt.figure(figsize=(12.2, 4.0), layout="constrained")
    grid = fig.add_gridspec(2, 1, height_ratios=[12, 1.2])
    ax = fig.add_subplot(grid[0])
    note = fig.add_subplot(grid[1])
    note.axis("off")
    picture = ax.imshow(np.ma.masked_invalid(values), cmap=cmap,
                        norm=colors.Normalize(vmin=0, vmax=1), aspect="auto")
    ax.set_xticks(range(15), zones)
    ax.set_yticks(range(3), list("ABC"))
    ax.set_xlabel("服务区")
    ax.set_ylabel("运输机型")
    ax.set_xticks(np.arange(-0.5, 15, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, 3, 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=1.4)
    ax.tick_params(which="minor", bottom=False, left=False)
    for row in range(3):
        for column in range(15):
            ratio = values[row, column]
            ax.text(column, row, labels[row][column], ha="center", va="center",
                    fontsize=8.7, fontweight="bold" if "*" in labels[row][column] else "normal",
                    color="white" if math.isfinite(ratio) and ratio > 0.66 else "#17242F")
    colorbar = fig.colorbar(picture, ax=ax, fraction=0.032, pad=0.018)
    colorbar.set_label("安全载荷 / 额定载荷")
    colorbar.set_ticks(np.linspace(0, 1, 6))
    colorbar.set_ticklabels([f"{int(v * 100)}%" for v in np.linspace(0, 1, 6)])
    note.text(0, 0.5, "单元格数字：最大安全载荷（kg）；* 表示受返航能量约束，未标星表示达到额定载荷上限。",
              ha="left", va="center", fontsize=9, color="#364451")
    _save_figure(fig, destination)
    return {"energy_limited_cells": limited, "rated_cells": 45 - len(limited) -
            int(np.isnan(values).sum()), "unreachable_cells": int(np.isnan(values).sum())}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def create_figures(result_dir: Path, output_dir: Path) -> dict[str, object]:
    data = load_results(result_dir)
    if output_dir.exists():
        raise FileExistsError(f"不覆盖已有图表目录：{output_dir}")
    font_name = configure_plot_style()
    output_dir.mkdir(parents=True)
    comparison = output_dir / "图1_两种可行基线对比"
    payload = output_dir / "图2_各区安全载荷边界"
    reductions = plot_comparison(data, comparison)
    payload_summary = plot_safe_payload(data, payload)
    artifact_names = [f"{stem.name}.{extension}" for stem in (comparison, payload)
                      for extension in ("pdf", "png")]
    for name in artifact_names:
        if (output_dir / name).stat().st_size < 10_000:
            raise ValueError(f"图表文件过小，可能导出失败：{name}")
    provenance = {
        "source_result_dir": str(result_dir.resolve()),
        "source_sha256": {name: _sha256(result_dir / name) for name in SOURCE_FILES},
        "source_qa_pass_count": data["qa_count"],
        "model_status": "feasible baseline, not optimized",
        "transformations": ["cumulative_operation_s / 3600 -> hours",
                            "percentage decrease = 100*(1-FFD/single_box)",
                            "heatmap ratio = q_max_kg / rated_payload_kg",
                            "no filtering or smoothing"],
        "comparison_reduction_pct": reductions,
        "payload_summary": payload_summary,
        "font": font_name,
        "matplotlib_version": matplotlib.__version__,
        "artifacts": artifact_names,
    }
    (output_dir / "图表数据来源.json").write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "图表说明.txt").write_text(
        "第一问目前已有两张结果图。\n"
        "图1比较单箱保底与同区FFD组批在架次、能耗和累计作业时间上的差异；\n"
        "图2展示A/B/C三种机型在15个服务区的最大安全载荷，星号代表受返航能量约束。\n"
        "这些是满足当前建模假设的可行基线，不是第一问最终最优解。\n"
        "PNG可直接发微信；PDF为矢量图，可用于论文排版。\n",
        encoding="utf-8")
    return provenance


def main() -> None:
    parser = argparse.ArgumentParser(description="D题第一问单点往返基线")
    parser.add_argument("--self-test", action="store_true", help="只运行内置测试，不生成正式结果")
    parser.add_argument("--data-run", type=Path, help="指定已验收数据运行目录；默认最新通过验收的一次")
    parser.add_argument("--output-dir", type=Path, help="指定新的结果目录；默认自动生成")
    parser.add_argument("--figure-dir", type=Path, help="指定新的图表目录；默认与表格目录同名时间戳")
    args = parser.parse_args()
    if args.self_test:
        raise SystemExit(pytest.main([str(Path(__file__).resolve()), "-q", "-p", "no:cacheprovider", "-o", "python_files=*.py"]))
    data_run = args.data_run.resolve() if args.data_run else find_latest_validated_run()
    result = run_model(data_run)
    output_dir = args.output_dir.resolve() if args.output_dir else (
        CODE_DIR / "1_outputs" / datetime.now().strftime("q1_base_%y%m%d_%H%M%S_table"))
    figure_dir = args.figure_dir.resolve() if args.figure_dir else figure_dir_for(output_dir)
    if figure_dir.exists():
        raise FileExistsError(f"不覆盖既有图表：{figure_dir}")
    save_results(result, output_dir)
    report = create_figures(output_dir, figure_dir)
    single, ffd = result["metrics"]
    print(f"Q1 baseline PASS: 45 safe loads, 80 boxes, {len(result['checks'])} checks")
    print(f"single-box: {single['sorties']} sorties, {single['total_energy_kwh']:.6f} kWh, {single['cumulative_operation_s']:.3f} s")
    print(f"FFD: {ffd['sorties']} sorties, {ffd['total_energy_kwh']:.6f} kWh, {ffd['cumulative_operation_s']:.3f} s")
    print(f"output: {output_dir}")
    print(f"figures: {figure_dir} ({len(report['artifacts'])} PDF/PNG)")




# 内置自检：仅通过 --self-test 运行，不影响正常求解。

def toy_inputs():
    model = {
        "type_id": "T", "empty_mass_kg": 5.0, "max_payload_kg": 10.0,
        "capacity_m3": 1.0, "cruise_speed_mps": 10.0,
        "range_empty_m": 1000.0, "range_full_m": 500.0,
        "energy_use_kwh": 2.0, "return_soc_min": 0.2,
        "prep_s": 30.0, "load_per_box_s": 2.0,
        "handoff_base_s": 5.0, "handoff_per_box_s": 1.0,
        "climb_speed_mps": 2.0, "descent_speed_mps": 2.0,
        "climb_efficiency": 0.72,
    }
    outward = {"distance_m": 300.0, "climb_m": 0.0, "descent_m": 0.0}
    homeward = dict(outward)
    box = {"box_id": "B01", "zone_id": "S001", "mass_kg": 3.0, "volume_m3": 0.1}
    return model, outward, homeward, box


def test_range_boundary_and_energy_limited_payload():
    model, outward, homeward, _ = toy_inputs()
    assert equivalent_range_m(model, 0.0) == 1000.0
    assert equivalent_range_m(model, 10.0) == 500.0
    result = max_safe_payload_kg(model, outward, homeward)
    assert result["status"] == "energy_limited"
    assert 0 < result["q_max_kg"] < 10
    assert result["energy_at_qmax_kwh"] <= 1.6 + 1e-8
    unreachable = dict(outward, distance_m=500.0)
    assert max_safe_payload_kg(model, unreachable, unreachable)["status"] == "unreachable"


def test_sortie_time_and_return_soc():
    model, outward, homeward, box = toy_inputs()
    result = sortie_metrics(model, outward, homeward, [box])
    assert result["feasible"]
    assert result["flight_s"] == pytest.approx(60.0)
    assert result["operation_s"] == pytest.approx(98.0)
    assert result["return_soc"] == pytest.approx(1 - result["energy_kwh"] / 2.0)
    assert result["energy_kwh"] <= 1.6
    assert result["energy_back_kwh"] != result["energy_out_kwh"]


def test_two_box_handoff_base_is_charged_once():
    model, outward, homeward, box = toy_inputs()
    second = dict(box, box_id="B02", mass_kg=2.0)
    result = sortie_metrics(model, outward, homeward, [box, second])
    assert result["handoff_s"] == pytest.approx(7.0)
    assert result["loading_s"] == pytest.approx(4.0)
    assert result["operation_s"] == pytest.approx(101.0)


def test_volume_is_independent_of_mass():
    model, outward, homeward, box = toy_inputs()
    oversized = dict(box, volume_m3=1.01)
    assert best_type_for_batch([oversized], {"T": model}, outward, homeward) is None


def test_choose_lower_energy_type_then_type_id():
    model, outward, homeward, box = toy_inputs()
    alternative = dict(model, type_id="U", energy_use_kwh=2.5)
    chosen = best_type_for_batch([box], {"T": model, "U": alternative}, outward, homeward)
    assert chosen is not None
    assert chosen["type_id"] == "T"
    equal = dict(model, type_id="A")
    chosen_equal = best_type_for_batch([box], {"T": model, "A": equal}, outward, homeward)
    assert chosen_equal is not None
    assert chosen_equal["type_id"] == "A"


def test_one_type_unreachable_does_not_invalidate_feasible_problem(monkeypatch):
    model, outward, homeward, _ = toy_inputs()
    boxes = [{"box_id": f"B{i:03d}", "zone_id": f"S{min(i // 5 + 1, 15):03d}",
              "mass_kg": 9.475, "volume_m3": 2.011 / 80}
             for i in range(80)]
    a = dict(model, type_id="A", return_soc_min=0.99)
    b = dict(model, type_id="B", max_payload_kg=100.0, energy_use_kwh=10.0,
             range_empty_m=10000.0, range_full_m=5000.0)
    c = dict(b, type_id="C")
    arcs = {}
    for zone in {box["zone_id"] for box in boxes}:
        arcs[("O01", zone)] = dict(outward, valid_flag=True)
        arcs[(zone, "O01")] = dict(homeward, valid_flag=True)
    monkeypatch.setattr(sys.modules[__name__], "_read_inputs", lambda _: (boxes, {"A": a, "B": b, "C": c}, arcs))
    result = run_model(Path("synthetic"))
    assert any(row["status"] == "unreachable" for row in result["max_payloads"])
    assert all(check["status"] == "PASS" for check in result["checks"])
    assert len(result["single_box"]) == 80


def test_official_data_baseline_replays_all_boxes():
    try:
        data_run = find_latest_validated_run()
    except FileNotFoundError:
        pytest.skip("尚无已验收的数据运行目录")
    result = run_model(data_run)
    assert len(result["max_payloads"]) == 45
    assert len(result["single_box"]) == 80
    assert sum(len(batch["box_ids"]) for batch in result["ffd"]) == 80
    assert len(result["ffd"]) <= len(result["single_box"])
    assert all(check["status"] == "PASS" for check in result["checks"])


def test_official_hand_calculation_anchors():
    try:
        data_run = find_latest_validated_run()
    except FileNotFoundError:
        pytest.skip("尚无已验收的数据运行目录")
    types = pd.read_csv(data_run / "clean" / "transport_types.csv").set_index("type_id")
    arcs = pd.read_csv(data_run / "derived" / "arc_geometry.csv").set_index(["from_id", "to_id"])
    a = types.loc["A"].to_dict()
    c = types.loc["C"].to_dict()
    a_out = arcs.loc[("O01", "S001")].to_dict()
    a_back = arcs.loc[("S001", "O01")].to_dict()
    a_result = max_safe_payload_kg(a, a_out, a_back)
    assert a_result["energy_empty_kwh"] == pytest.approx(1.161539158, abs=1e-7)
    assert a_result["q_max_kg"] == 25.0
    assert a_result["energy_at_qmax_kwh"] == pytest.approx(1.312722009, abs=1e-7)
    c_result = max_safe_payload_kg(c, arcs.loc[("O01", "S008")].to_dict(),
                                   arcs.loc[("S008", "O01")].to_dict())
    assert c_result["q_max_kg"] == pytest.approx(58.903108, abs=1e-6)
    assert c_result["energy_at_qmax_kwh"] <= c_result["energy_limit_kwh"]


def test_export_reopens_csv_before_ready(tmp_path):
    try:
        data_run = find_latest_validated_run()
    except FileNotFoundError:
        pytest.skip("尚无已验收的数据运行目录")
    result = run_model(data_run)
    destination = tmp_path / "q1_0_baseline"
    save_results(result, destination)
    assert (destination / "Q1_READY.txt").exists()
    assert list(pd.read_csv(destination / "Q1_单点组批.csv").columns) == TEMPLATE_COLUMNS


if __name__ == "__main__":
    main()
